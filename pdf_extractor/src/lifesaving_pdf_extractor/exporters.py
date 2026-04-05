from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from .models import ExtractedRecord, ReviewFlag

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:  # pragma: no cover - dependency is optional for unit tests
    Workbook = None
    Font = None


def _discipline_fields(prefix: str) -> tuple[str, str, str, str, str]:
    return (
        f"{prefix}_place",
        f"{prefix}_time",
        f"{prefix}_points",
        f"{prefix}_code",
        f"{prefix}_penalty",
    )


DISCIPLINE_FIELDS: dict[str, tuple[str, str, str, str, str]] = {
    "event_100m_lifesaver": _discipline_fields("event_100m_lifesaver"),
    "event_50m_manikin_carry": _discipline_fields("event_50m_manikin_carry"),
    "event_200m_super_lifesaver": _discipline_fields("event_200m_super_lifesaver"),
    "event_100m_rescue_medley": _discipline_fields("event_100m_rescue_medley"),
    "event_100m_manikin_carry_with_fins": _discipline_fields("event_100m_manikin_carry_with_fins"),
    "event_200m_obstacle": _discipline_fields("event_200m_obstacle"),
    "event_100m_obstacle": _discipline_fields("event_100m_obstacle"),
    "event_50m_manikin_carry_with_fins": _discipline_fields("event_50m_manikin_carry_with_fins"),
    "event_50m_finswim": _discipline_fields("event_50m_finswim"),
    "event_50m_obstacle": _discipline_fields("event_50m_obstacle"),
    "event_50m_combined_swim": _discipline_fields("event_50m_combined_swim"),
    "event_50m_freestyle": _discipline_fields("event_50m_freestyle"),
    "event_25m_freestyle": _discipline_fields("event_25m_freestyle"),
    "event_25m_breaststroke": _discipline_fields("event_25m_breaststroke"),
    "event_25m_backstroke_no_arms": _discipline_fields("event_25m_backstroke_no_arms"),
    "event_25m_manikin": _discipline_fields("event_25m_manikin"),
}

FINAL_CSV_FIELDS = (
    "name",
    "gender",
    "altersklasse",
    "jahrgang",
    "gliederung_ortsgruppe",
    "bezirk_q_gliederung",
    "reserved_07",
    "athlete_country",
    "gesamtplatzierung",
    "mehrkampf_punkte",
    "reserved_11",
    "reserved_12",
    "reserved_13",
    *DISCIPLINE_FIELDS["event_100m_lifesaver"],
    *DISCIPLINE_FIELDS["event_50m_manikin_carry"],
    *DISCIPLINE_FIELDS["event_200m_super_lifesaver"],
    *DISCIPLINE_FIELDS["event_100m_rescue_medley"],
    *DISCIPLINE_FIELDS["event_100m_manikin_carry_with_fins"],
    *DISCIPLINE_FIELDS["event_200m_obstacle"],
    *DISCIPLINE_FIELDS["event_100m_obstacle"],
    "reserved_49",
    "reserved_50",
    "reserved_51",
    "reserved_52",
    "reserved_53",
    *DISCIPLINE_FIELDS["event_50m_manikin_carry_with_fins"],
    *DISCIPLINE_FIELDS["event_50m_finswim"],
    *DISCIPLINE_FIELDS["event_50m_obstacle"],
    *DISCIPLINE_FIELDS["event_50m_combined_swim"],
    *DISCIPLINE_FIELDS["event_50m_freestyle"],
    *DISCIPLINE_FIELDS["event_25m_freestyle"],
    *DISCIPLINE_FIELDS["event_25m_breaststroke"],
    *DISCIPLINE_FIELDS["event_25m_backstroke_no_arms"],
    *DISCIPLINE_FIELDS["event_25m_manikin"],
    "reserved_99",
    "datum_excel",
    "reserved_101",
    "bahnlaenge",
    "reserved_103",
    "austragungsland",
    "austragungsort",
    "reserved_106",
    "runde",
    "wertung",
    "wettkampfname",
)

SUMMARY_ROW_PATTERNS = (
    re.compile(
        r"^(?P<place>\d+|-)\s+(?P<name>.+?)\s{2,}(?P<club>.+?)\s{2,}"
        r"(?P<qualifier>.+?)\s{2,}(?P<year>\d{2,4})\s+"
        r"(?P<points>\.\.\.|n\.a\.|[\d.,]+)\s+(?P<diff>[-\d.,]+)\s*(?P<tail>.*)$",
        flags=re.IGNORECASE,
    ),
    re.compile(
        r"^(?P<place>\d+|-)\s+(?P<name>.+?)\s{2,}(?P<club>.+?)\s{2,}"
        r"(?P<year>\d{2,4})\s+(?P<points>\.\.\.|n\.a\.|[\d.,]+)\s+"
        r"(?P<diff>[-\d.,]+)\s*(?P<tail>.*)$",
        flags=re.IGNORECASE,
    ),
)


class ExcelExporter:
    def export(self, records: list[ExtractedRecord], output_path: Path, provisional_fields: tuple[str, ...]) -> Path:
        del provisional_fields
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if Workbook is None:
            raise RuntimeError("openpyxl is not installed. Install dependencies from requirements.txt first.")

        if not records:
            empty_path = output_path.parent / "empty_export.xlsx"
            self._write_rows(empty_path, [])
            return empty_path

        grouped: dict[Path, list[ExtractedRecord]] = defaultdict(list)
        for record in records:
            grouped[record.source_file].append(record)

        written_paths: list[Path] = []
        for source_file, source_records in sorted(grouped.items(), key=lambda item: str(item[0]).casefold()):
            workbook_path = output_path.parent / f"{self._safe_stem(source_file.stem)}.xlsx"
            rows = self._build_rows(source_records)
            self._write_rows(workbook_path, rows)
            written_paths.append(workbook_path)

        return written_paths[0] if len(written_paths) == 1 else output_path.parent

    def _write_rows(self, output_path: Path, rows: list[dict[str, str]]) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Ergebnisse"
        worksheet.append(list(FINAL_CSV_FIELDS))

        if Font is not None:
            for cell in worksheet[1]:
                cell.font = Font(bold=True)

        for row in rows:
            worksheet.append([self._excel_cell_value(field, row.get(field, "")) for field in FINAL_CSV_FIELDS])

        workbook.save(output_path)

    def _build_rows(self, records: list[ExtractedRecord]) -> list[dict[str, str]]:
        rows: list[dict[str, str]] = []
        summary_rows: dict[tuple[str, str, str, str], dict[str, str]] = {}

        ordered_records = sorted(
            records,
            key=lambda record: (
                0 if self._is_summary_record(record) else 1,
                record.page_number or 0,
                self._canonical(record.data.get("name", "")),
                self._canonical(record.data.get("disziplin", "")),
            ),
        )

        for record in ordered_records:
            if not self._is_summary_record(record):
                continue

            row = self._build_base_row(record)
            row["gesamtplatzierung"] = self._clean_place(record.data.get("platz", ""))
            row["mehrkampf_punkte"] = self._clean_decimal(record.data.get("punkte", ""))
            self._populate_summary_disciplines(row, record)
            row["wertung"] = self._derive_summary_wertung_from_row(row) or self._derive_wertung(record)
            rows.append(row)
            summary_rows[self._row_key(record)] = row

        for record in ordered_records:
            if self._is_summary_record(record):
                continue

            if self._is_status_only_record(record):
                target_row = summary_rows.get(self._row_key(record))
                if target_row is not None and self._apply_status_to_existing_row(target_row, record):
                    continue

            row = self._build_base_row(record)
            row["wertung"] = self._derive_wertung(record)
            self._populate_single_discipline(row, record)
            rows.append(row)

        return rows

    def _build_base_row(self, record: ExtractedRecord) -> dict[str, str]:
        row = {field: "" for field in FINAL_CSV_FIELDS}
        row["name"] = record.data.get("name", "")
        row["gender"] = self._gender_code(record.data.get("geschlecht", ""))
        row["altersklasse"] = self._format_age_group(record.data.get("altersklasse", ""))
        row["jahrgang"] = self._two_digit_year(record.data.get("jahrgang", ""))
        row["gliederung_ortsgruppe"] = record.data.get("verein", "")
        row["bezirk_q_gliederung"] = self._extract_region(record)
        row["athlete_country"] = "GER"
        row["datum_excel"] = self._excel_date(record.data.get("datum", ""))
        row["bahnlaenge"] = self._extract_pool_length(record)
        row["austragungsland"] = self._extract_country(record)
        row["austragungsort"] = self._extract_city(record.data.get("ort", ""))
        row["runde"] = self._derive_round(record)
        row["wettkampfname"] = record.data.get("wettkampfname", "")
        return row

    def _populate_summary_disciplines(self, row: dict[str, str], record: ExtractedRecord) -> None:
        discipline_keys = self._extract_summary_discipline_keys(record)
        values = self._extract_summary_discipline_values(record.raw_excerpt, len(discipline_keys))
        if values and len(values) < len(discipline_keys):
            discipline_keys = discipline_keys[: len(values)]

        for index, key in enumerate(discipline_keys):
            if index >= len(values):
                continue

            discipline_value = values[index]
            time_value = discipline_value.get("time", "")
            points_value = discipline_value.get("points", "")
            code = discipline_value.get("code", "")
            penalty = discipline_value.get("penalty", "")

            self._write_discipline_block(
                row=row,
                discipline_key=key,
                place="",
                time_value=time_value,
                points_value=points_value,
                code=code,
                penalty=penalty,
            )

    def _populate_single_discipline(self, row: dict[str, str], record: ExtractedRecord) -> None:
        discipline_key = self._discipline_key(record.data.get("disziplin", ""))
        if discipline_key is None:
            return

        code, penalty = self._extract_code_and_penalty(record)
        self._write_discipline_block(
            row=row,
            discipline_key=discipline_key,
            place=self._clean_place(record.data.get("platz", "")),
            time_value=record.data.get("zeit", ""),
            points_value=record.data.get("punkte", ""),
            code=code,
            penalty=penalty,
        )

    def _apply_status_to_existing_row(self, row: dict[str, str], record: ExtractedRecord) -> bool:
        discipline_key = self._discipline_key(record.data.get("disziplin", ""))
        if discipline_key is None:
            return False

        code, penalty = self._extract_code_and_penalty(record)
        fields = DISCIPLINE_FIELDS[discipline_key]
        if code:
            row[fields[3]] = code
        if penalty:
            row[fields[4]] = penalty
        return True

    def _write_discipline_block(
        self,
        row: dict[str, str],
        discipline_key: str,
        place: str,
        time_value: str,
        points_value: str,
        code: str,
        penalty: str,
    ) -> None:
        place_field, time_field, points_field, code_field, penalty_field = DISCIPLINE_FIELDS[discipline_key]
        row[place_field] = place
        row[time_field] = self._format_time(time_value)
        row[points_field] = self._clean_decimal(points_value)
        row[code_field] = code
        row[penalty_field] = penalty

    @staticmethod
    def _safe_stem(value: str) -> str:
        cleaned = re.sub(r"[^\w.-]+", "_", value, flags=re.ASCII).strip("_")
        return cleaned or "protocol"

    @staticmethod
    def _excel_cell_value(field_name: str, value: str) -> str | int:
        if field_name == "datum_excel" and value.isdigit():
            return int(value)
        return value

    @staticmethod
    def _is_summary_record(record: ExtractedRecord) -> bool:
        return record.data.get("disziplin", "") == "Einzelwertung"

    @staticmethod
    def _is_status_only_record(record: ExtractedRecord) -> bool:
        return not record.data.get("platz", "") and not record.data.get("zeit", "") and not record.data.get("punkte", "")

    def _row_key(self, record: ExtractedRecord) -> tuple[str, str, str, str]:
        return (
            self._canonical(record.data.get("name", "")),
            self._canonical(record.data.get("geschlecht", "")),
            self._canonical(record.data.get("altersklasse", "")),
            self._canonical(record.data.get("verein", "")),
        )

    def _extract_region(self, record: ExtractedRecord) -> str:
        for pattern in (r"(?:Q-Gld|LV)=([^;]+)",):
            match = re.search(pattern, record.data.get("bemerkung", ""), flags=re.IGNORECASE)
            if match:
                candidate = match.group(1).strip()
                if self._looks_like_region(candidate):
                    return candidate
        return record.data.get("bezirk_q_gliederung", "")

    def _extract_summary_discipline_names(self, record: ExtractedRecord) -> list[str]:
        raw_disciplines = record.data.get("summary_disciplines", "")
        if raw_disciplines:
            return [item.strip() for item in raw_disciplines.split("|") if item.strip()]

        match = re.search(r"Disziplinen=(.+)$", record.data.get("bemerkung", ""))
        if not match:
            return []
        return [item.strip() for item in match.group(1).split("|") if item.strip()]

    def _extract_summary_discipline_keys(self, record: ExtractedRecord) -> list[str]:
        raw_keys = record.data.get("summary_discipline_keys", "")
        if raw_keys:
            return [item.strip() for item in raw_keys.split("|") if item.strip()]

        keys: list[str] = []
        for name in self._extract_summary_discipline_names(record):
            key = self._discipline_key(name)
            if key is None:
                continue
            keys.append(key)
        return keys

    def _extract_summary_discipline_values(self, raw_excerpt: str, expected_count: int) -> list[dict[str, str]]:
        tail = self._extract_summary_tail(raw_excerpt)
        if not tail:
            return []

        structured_values = self._extract_structured_summary_discipline_values(tail, expected_count)
        if structured_values:
            return structured_values

        matches = re.findall(
            r"([0-9:.,]+|n\.a\.|DNS|DQ|DNF)\s+([0-9.,]+|n\.a\.|DNS|DQ|DNF)",
            tail,
            flags=re.IGNORECASE,
        )
        return [
            {
                "time": "" if self._is_status_token(time_value.strip()) else time_value.strip(),
                "points": "" if self._is_status_token(points_value.strip()) else points_value.strip(),
                "code": self._code_from_token(time_value.strip()) if self._is_status_token(time_value.strip()) else "",
                "penalty": self._penalty_from_token(time_value.strip()) if self._is_status_token(time_value.strip()) else "",
            }
            for time_value, points_value in matches
        ]

    def _extract_summary_tail(self, raw_excerpt: str) -> str:
        normalized_excerpt = self._normalize_text_layer(raw_excerpt, collapse_spaces=False)
        for pattern in SUMMARY_ROW_PATTERNS:
            match = pattern.match(normalized_excerpt)
            if match:
                return match.group("tail").strip()
        return ""

    def _extract_structured_summary_discipline_values(self, tail: str, expected_count: int) -> list[dict[str, str]]:
        tokens = self._tokenize_summary_tail(tail)
        if not tokens:
            return []
        start_indices = self._summary_slot_start_indices(tokens)
        if not start_indices:
            return []

        values: list[dict[str, str]] = []
        remaining_empty_slots = max(0, expected_count - len(start_indices))

        leading_placeholders = start_indices[0]
        while leading_placeholders > 0 and remaining_empty_slots > 0:
            values.append(self._new_summary_value(""))
            leading_placeholders -= 1
            remaining_empty_slots -= 1

        for position, start_index in enumerate(start_indices):
            next_start = start_indices[position + 1] if position + 1 < len(start_indices) else len(tokens)
            slot_tokens = tokens[start_index:next_start]
            values.append(self._parse_summary_slot_tokens(slot_tokens))

            trailing_placeholders = self._count_trailing_placeholders(slot_tokens)
            candidate_empty_slots = max(0, trailing_placeholders - 1)
            while candidate_empty_slots > 0 and remaining_empty_slots > 0:
                values.append(self._new_summary_value(""))
                candidate_empty_slots -= 1
                remaining_empty_slots -= 1

        if expected_count > 0 and len(values) < expected_count:
            values.extend(self._new_summary_value("") for _ in range(expected_count - len(values)))

        return values[:expected_count] if expected_count > 0 else values

    def _tokenize_summary_tail(self, tail: str) -> list[str]:
        normalized = self._normalize_text_layer(tail)
        normalized = normalized.replace("...", " ... ")
        normalized = re.sub(r"\s+", " ", normalized).strip()
        return [token for token in normalized.split(" ") if token]

    def _new_summary_value(self, anchor: str) -> dict[str, str]:
        value = {"time": "", "points": "", "code": "", "penalty": ""}
        stripped_anchor = anchor.strip()

        if not stripped_anchor:
            return value

        if self._looks_like_time_value(stripped_anchor):
            value["time"] = stripped_anchor
            return value

        if self._is_status_token(stripped_anchor):
            value["code"] = self._code_from_token(stripped_anchor)
            value["penalty"] = self._penalty_from_token(stripped_anchor)
            return value

        return value

    def _summary_slot_start_indices(self, tokens: list[str]) -> list[int]:
        starts: list[int] = []
        for index, token in enumerate(tokens):
            cleaned = token.rstrip(":;")
            if cleaned == "...":
                continue

            if index > 0 and self._looks_like_summary_number(tokens[index - 1].rstrip(":;")) and self._is_status_token(cleaned):
                continue

            if self._looks_like_summary_number(cleaned):
                next_token = tokens[index + 1].rstrip(":;") if index + 1 < len(tokens) else ""
                if self._is_status_token(next_token):
                    starts.append(index)
                continue

            if self._looks_like_summary_anchor(cleaned):
                starts.append(index)

        return starts

    def _parse_summary_slot_tokens(self, slot_tokens: list[str]) -> dict[str, str]:
        value = self._new_summary_value("")
        content = [token.rstrip(":;") for token in slot_tokens if token.rstrip(":;") != "..."]
        if not content:
            return value

        index = 0
        if len(content) >= 2 and self._looks_like_summary_number(content[0]) and self._is_status_token(content[1]):
            value["points"] = content[0]
            value["code"] = self._code_from_token(content[1])
            value["penalty"] = self._penalty_from_token(content[1])
            index = 2
        elif self._looks_like_time_value(content[0]):
            value["time"] = content[0]
            index = 1
        elif self._is_status_token(content[0]):
            value["code"] = self._code_from_token(content[0])
            value["penalty"] = self._penalty_from_token(content[0])
            index = 1

        for token in content[index:]:
            lowered = token.casefold()
            if self._looks_like_summary_code(token):
                value["code"] = token.upper()
                continue
            if lowered.startswith("disq"):
                value["penalty"] = "disq."
                if not value["code"]:
                    value["code"] = "DQ"
                continue
            if self._looks_like_summary_number(token):
                if not value["points"]:
                    value["points"] = token
                    continue
                if token in {"50", "100", "150", "200"} and not value["penalty"]:
                    value["penalty"] = token
                continue
            if self._is_status_token(token) and not value["code"]:
                value["code"] = self._code_from_token(token)
                value["penalty"] = self._penalty_from_token(token)

        return value

    @staticmethod
    def _count_trailing_placeholders(slot_tokens: list[str]) -> int:
        count = 0
        for token in reversed(slot_tokens):
            if token.rstrip(":;") != "...":
                break
            count += 1
        return count

    def _looks_like_summary_anchor(self, token: str) -> bool:
        if self._looks_like_time_value(token):
            return True
        return self._is_status_token(token)

    @staticmethod
    def _looks_like_time_value(token: str) -> bool:
        return bool(
            re.fullmatch(r"\d{1,2}:\d{2}[.,]\d{1,2}", token)
            or re.fullmatch(r"\d{2}:\d{2}:\d{2}\.\d{2}", token)
            or re.fullmatch(r"\d{1,2}[.,]\d{1,2}", token)
        )

    @staticmethod
    def _looks_like_summary_code(token: str) -> bool:
        return bool(re.fullmatch(r"[A-Z]{1,3}\d{1,2}", token, flags=re.IGNORECASE))

    @staticmethod
    def _looks_like_summary_number(token: str) -> bool:
        return bool(re.fullmatch(r"\d+(?:[.,]\d+)?", token))

    def _derive_wertung(self, record: ExtractedRecord) -> str:
        if self._is_summary_record(record):
            explicit_count = record.data.get("summary_discipline_count", "").strip()
            if explicit_count.isdigit():
                count = int(explicit_count)
                return f"{count}-Kampf" if count > 0 else ""
            discipline_keys = self._extract_summary_discipline_keys(record)
            return f"{len(discipline_keys)}-Kampf" if discipline_keys else ""
        if self._is_status_only_record(record):
            return ""
        return "Einzelkampf"

    def _derive_summary_wertung_from_row(self, row: dict[str, str]) -> str:
        count = 0
        for fields in DISCIPLINE_FIELDS.values():
            if any(row.get(field, "").strip() for field in fields):
                count += 1
        return f"{count}-Kampf" if count > 0 else ""

    def _derive_round(self, record: ExtractedRecord) -> str:
        round_name = self._extract_round_name(record)
        if not round_name:
            return "1" if record.data.get("disziplin", "") else ""
        canonical = self._canonical(round_name)
        return "2" if "finale" in canonical or "endlauf" in canonical else "1"

    @staticmethod
    def _extract_round_name(record: ExtractedRecord) -> str:
        match = re.search(r"Runde=([^;]+)", record.data.get("bemerkung", ""))
        return match.group(1).strip() if match else ""

    def _extract_code_and_penalty(self, record: ExtractedRecord) -> tuple[str, str]:
        combined = " ".join(
            part
            for part in (
                record.data.get("dq_status", ""),
                record.data.get("bemerkung", ""),
                record.raw_excerpt,
            )
            if part
        )
        dq_status = record.data.get("dq_status", "")

        code_match = re.search(r"\b([A-Z]\d+)\b", combined)
        if code_match:
            code = code_match.group(1).upper()
        elif dq_status == "DNS":
            code = "n.a." if "nicht angetreten" in self._canonical(combined) or "n.a." in combined.casefold() else "DNS"
        elif dq_status == "DQ":
            code = "DQ"
        elif dq_status == "DNF":
            code = "DNF"
        else:
            code = ""

        penalty_match = re.search(r"\b(50|100|150|200)\b", combined)
        if penalty_match:
            penalty = penalty_match.group(1)
        elif dq_status == "DQ" or "disqual" in self._canonical(combined):
            penalty = "disq."
        else:
            penalty = ""

        return code, penalty

    def _discipline_key(self, discipline_name: str) -> str | None:
        canonical = self._canonical(discipline_name)
        canonical = canonical.replace("obsatcle", "obstacle")
        canonical = canonical.replace("manakin", "manikin")

        if "100m retten mit gurt" in canonical or "100m lifesaver" in canonical or "manikin tow with fins" in canonical:
            return "event_100m_lifesaver"
        if "50m retten mit flossen" in canonical:
            return "event_50m_manikin_carry_with_fins"
        if "50m retten" in canonical or "50m manikin carry" in canonical:
            return "event_50m_manikin_carry"
        if "200m super-lifesaver" in canonical or "200m super lifesaver" in canonical or "superlifesaver" in canonical:
            return "event_200m_super_lifesaver"
        if "100m kombi" in canonical or "kombinierte rettungsuebung" in canonical or "rescue medley" in canonical:
            return "event_100m_rescue_medley"
        if "100m retten mit flossen" in canonical or canonical == "100m retten" or "100m manikin carry with fins" in canonical:
            return "event_100m_manikin_carry_with_fins"
        if "200m hindernis" in canonical or "200m obstacle swim" in canonical:
            return "event_200m_obstacle"
        if "100m hindernis" in canonical or "100m obstacle swim" in canonical:
            return "event_100m_obstacle"
        if "50m flossen" in canonical or "flossenschwimmen" in canonical:
            return "event_50m_finswim"
        if "50m hindernis" in canonical or "50m hindernisschwimmen" in canonical:
            return "event_50m_obstacle"
        if "50m k. schwimmen" in canonical or "50m kombiniertes schwimmen" in canonical:
            return "event_50m_combined_swim"
        if "50m freistil" in canonical or "50m kraul" in canonical:
            return "event_50m_freestyle"
        if "25m freistil" in canonical or "25m kraul" in canonical:
            return "event_25m_freestyle"
        if "25m brust" in canonical:
            return "event_25m_breaststroke"
        if "25m rue." in canonical or "25m ruecken" in canonical or ("25m" in canonical and "ohne arm" in canonical):
            return "event_25m_backstroke_no_arms"
        if "25m puppe" in canonical:
            return "event_25m_manikin"
        return None

    @staticmethod
    def _gender_code(value: str) -> str:
        lowered = value.casefold()
        if lowered.startswith("w"):
            return "w"
        if lowered.startswith("m"):
            return "m"
        return ""

    @staticmethod
    def _two_digit_year(value: str) -> str:
        digits = re.sub(r"\D+", "", value)
        if len(digits) >= 2:
            return digits[-2:]
        return ""

    @staticmethod
    def _format_age_group(value: str) -> str:
        stripped = value.strip()
        if not stripped:
            return ""
        stripped = re.sub(r"^(?:AK|Altersklasse)\s+", "", stripped, flags=re.IGNORECASE)
        return stripped.strip(" -")

    @staticmethod
    def _clean_place(value: str) -> str:
        return "" if value.strip() in {"", "-"} else value.strip()

    @staticmethod
    def _clean_decimal(value: str) -> str:
        stripped = value.strip()
        if not stripped or stripped in {"...", "n.a."}:
            return ""
        return stripped

    def _format_time(self, value: str) -> str:
        stripped = value.strip()
        if not stripped:
            return ""
        if self._is_status_token(stripped):
            return ""

        match = re.fullmatch(r"(?P<hours>\d{2}):(?P<minutes>\d{2}):(?P<seconds>\d{2})\.(?P<fraction>\d{2})", stripped)
        if match:
            total_minutes = int(match.group("hours")) * 60 + int(match.group("minutes"))
            return f"{total_minutes}:{match.group('seconds')},{match.group('fraction')}"

        match = re.fullmatch(r"(?P<minutes>\d{1,2}):(?P<seconds>\d{2})[.,](?P<fraction>\d{1,2})", stripped)
        if match:
            fraction = match.group("fraction").ljust(2, "0")[:2]
            return f"{int(match.group('minutes'))}:{match.group('seconds')},{fraction}"

        match = re.fullmatch(r"(?P<seconds>\d{1,2})[.,](?P<fraction>\d{1,2})", stripped)
        if match:
            fraction = match.group("fraction").ljust(2, "0")[:2]
            return f"0:{int(match.group('seconds')):02d},{fraction}"

        return stripped

    @staticmethod
    def _is_status_token(value: str) -> bool:
        return value.casefold() in {"n.a.", "dns", "dq", "dnf"}

    @staticmethod
    def _code_from_token(value: str) -> str:
        token = value.casefold()
        if token == "n.a.":
            return "n.a."
        return value.upper()

    @staticmethod
    def _penalty_from_token(value: str) -> str:
        return "disq." if value.casefold() == "dq" else ""

    def _excel_date(self, value: str) -> str:
        parsed_date = self._parse_date(value)
        if parsed_date is None:
            return ""
        return str((parsed_date - date(1899, 12, 30)).days)

    @staticmethod
    def _parse_date(value: str) -> date | None:
        stripped = value.strip()
        if not stripped:
            return None

        for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(stripped, pattern).date()
            except ValueError:
                continue
        return None

    def _extract_pool_length(self, record: ExtractedRecord) -> str:
        searchable = " ".join(
            item
            for item in (
                record.data.get("wettkampfname", ""),
                record.data.get("ort", ""),
                record.data.get("bemerkung", ""),
                record.raw_excerpt,
            )
            if item
        )
        canonical = self._canonical(searchable)
        if re.search(r"\b25\s*m\b.*\b(bahn|becken|pool)\b", canonical):
            return "25m"
        if re.search(r"\b50\s*m\b.*\b(bahn|becken|pool)\b", canonical):
            return "50m"
        return ""

    def _extract_country(self, record: ExtractedRecord) -> str:
        searchable = " ".join(item for item in (record.data.get("ort", ""), record.data.get("wettkampfname", "")) if item)
        canonical = self._canonical(searchable)
        country_map = {
            "deutschland": "GER",
            "germany": "GER",
            "frankreich": "Frankreich",
            "france": "Frankreich",
            "spanien": "Spanien",
            "spain": "Spanien",
            "oesterreich": "Oesterreich",
            "austria": "Oesterreich",
            "schweiz": "Schweiz",
            "switzerland": "Schweiz",
            "belgien": "Belgien",
            "belgium": "Belgien",
            "niederlande": "Niederlande",
            "netherlands": "Niederlande",
        }
        for token, export_value in country_map.items():
            if token in canonical:
                return export_value
        return ""

    @staticmethod
    def _extract_city(value: str) -> str:
        if not value.strip():
            return ""
        city = re.sub(r"\s*\([^)]*\)", "", value).strip()
        return city.split(",")[0].strip()

    @staticmethod
    def _looks_like_region(value: str) -> bool:
        candidate = value.strip()
        if not candidate:
            return False
        if len(candidate) > 30:
            return False
        if re.search(r"\d{1,2}:\d{2}", candidate):
            return False
        if re.search(r"\d{3,},\d{2}", candidate):
            return False
        if "..." in candidate:
            return False
        return True

    @staticmethod
    def _canonical(value: str) -> str:
        normalized = (
            value.replace("\u00e4", "ae")
            .replace("\u00f6", "oe")
            .replace("\u00fc", "ue")
            .replace("\u00c4", "ae")
            .replace("\u00d6", "oe")
            .replace("\u00dc", "ue")
            .replace("\u00df", "ss")
            .lower()
        )
        return re.sub(r"\s+", " ", normalized).strip()

    @staticmethod
    def _normalize_text_layer(value: str, collapse_spaces: bool = True) -> str:
        repaired = ExcelExporter._repair_mojibake(value)
        repaired = repaired.replace("\u00a0", " ").replace("\t", " ").replace("\ufeff", " ")
        repaired = repaired.replace("\u00c2", " ")
        repaired = re.sub(r"(?<=[A-Z])(\d+):", r"\1", repaired, flags=re.IGNORECASE)
        if collapse_spaces:
            repaired = re.sub(r"\s+", " ", repaired)
        else:
            repaired = re.sub(r"[ ]{3,}", "  ", repaired)
        return repaired.strip()

    @staticmethod
    def _repair_mojibake(value: str) -> str:
        if not value or ("\u00c3" not in value and "\u00c2" not in value):
            return value
        try:
            repaired = value.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            return value
        return repaired if repaired else value


class JSONExporter:
    def export(self, records: list[ExtractedRecord], output_path: Path, provisional_fields: tuple[str, ...]) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        payload = [record.to_json_dict(provisional_fields) for record in records]
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return output_path


class ReviewExporter:
    fieldnames = [
        "source_file",
        "page_number",
        "code",
        "message",
        "confidence",
        "parser_name",
        "status",
        "raw_excerpt",
    ]

    def export(self, records: list[ExtractedRecord], flags: list[ReviewFlag], output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        rows = self._collect_rows(records, flags)

        with output_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.fieldnames)
            writer.writeheader()
            for row in rows:
                writer.writerow(row)

        return output_path

    def _collect_rows(self, records: list[ExtractedRecord], flags: list[ReviewFlag]) -> list[dict[str, str]]:
        rows = [flag.to_row() for flag in flags]

        for record in records:
            for note in record.review_notes:
                rows.append(
                    {
                        "source_file": str(record.source_file),
                        "page_number": "" if record.page_number is None else str(record.page_number),
                        "code": "record_review_note",
                        "message": note,
                        "confidence": f"{record.confidence:.2f}",
                        "parser_name": record.parser_name,
                        "status": record.status.value,
                        "raw_excerpt": record.raw_excerpt,
                    }
                )

        return rows
